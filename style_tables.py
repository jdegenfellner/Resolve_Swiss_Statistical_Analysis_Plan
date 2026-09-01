#!/usr/bin/env python3
"""Gibt den Tabellen im docx sichtbare Linien.

Pandoc weist ihnen die Formatvorlage 'Table' zu, die es in der erzeugten
styles.xml aber nicht gibt. Word zeichnet sie deshalb ohne jeden Rand. Hier
werden Linien im Booktabs-Sinn gesetzt: eine oben, eine unter der Kopfzeile,
eine unten, keine senkrechten. Die Kopfzeile wird fett und wiederholt sich
ueber Seitenumbrueche.

Aufruf:  python3 style_tables.py <datei.docx>
"""
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def border(el, edge, size, val="single"):
    b = OxmlElement("w:" + edge)
    b.set(qn("w:val"), val)
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "0")
    b.set(qn("w:color"), "000000")
    el.append(b)


def style_table(t):
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    bd = OxmlElement("w:tblBorders")
    border(bd, "top", 12)
    border(bd, "bottom", 12)
    for edge in ("left", "right", "insideV", "insideH"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        bd.append(e)
    tblPr.append(bd)

    head = t.rows[0]
    trPr = head._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader")
    trPr.append(rep)
    for c in head.cells:
        tcPr = c._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:tcBorders")):
            tcPr.remove(old)
        tb = OxmlElement("w:tcBorders")
        border(tb, "bottom", 8)
        tcPr.append(tb)
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True


def main(path):
    doc = Document(path)
    n = 0
    for t in doc.tables:
        if len(t.rows) < 2:
            continue
        style_table(t)
        n += 1
    doc.save(path)
    print(f"{path}: {n} Tabellen mit Linien versehen")


if __name__ == "__main__":
    main(sys.argv[1])
